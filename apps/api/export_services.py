"""Export row builders and workbook writers for the Laravel-compatible API.

The Laravel application uses ``maatwebsite/excel`` for the XLSX endpoints.  A
small service module keeps the Django Ninja views thin while retaining the
original column order, labels and value transformations.  All data access is
through the Django ORM; no database-specific SQL is required.
"""

from __future__ import annotations

import io
import json
from collections import defaultdict
from typing import Iterable, Mapping, Optional, Sequence

from django.http import HttpResponse

from apps.core.models import Draw, Files, Person, Report, ReportPerson, User


def _cjk_pdf_font():
    """Reportlab's built-in Adobe CJK face; no external font file required.

    Non-embedded (viewer supplies the glyphs); registration is idempotent.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    name = "STSong-Light"
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(name))
    return name


REPORT_DATA_HEADINGS = [
    "所属单位", "乐团名称", "自选曲目名称", "指定曲目名称", "参报代码",
    "参展学校名称", "描述", "组别", "领队姓名", "领队电话", "联系地址或邮箱",
    "节目时长", "集体照", "视频文件", "状态", "正式队员", "预备队员", "指挥", "指导教师",
]

# data1/data2 的列集与《附件2 报名信息表》（0921 定稿通知 docx 附录）逐栏对齐：
# 附录栏目在前（栏名与栏序即表格原文），管理辅助列（报名学校/乐团名称/时长/
# 联系地址/简介/状态）保留在后。乐器 17 栏即附录「正式队员名单」的乐器槽。
ADMIN_DATA1_HEADINGS = [
    "序号", "参展学校名称", "领队姓名", "领队电话", "指挥", "指挥电话",
    "指导老师1", "指导老师1电话", "指导老师2", "指导老师2电话",
    "乐团类别", "参展组别", "指定曲目", "自选曲目", "参展人数",
    "正式队员名单（按乐器）", "预备队员名单", "备注", "用餐预约",
    "报名学校", "乐团名称", "节目时长", "联系地址", "乐团简介", "状态",
]

# The instrument columns of the data2 heading row are derived from this single
# list.  They were once written out twice and drifted: the heading row kept
# truncated spellings ("低音单簧", "低音大提") while the tally keys were full names.
INSTRUMENTS = (
    "短笛", "长笛", "单簧管", "低音单簧管", "中音萨克斯", "次中音萨克斯",
    "上低音萨克斯", "双簧管", "大管", "小号", "长号", "圆号", "上低音号", "大号",
    "打击乐", "低音大提琴", "其他",
)

ADMIN_DATA2_HEADINGS = [
    "序号", "报名学校", "参展学校名称", "乐团名称", "领队姓名", "领队电话",
    "指挥", "指挥电话", "指导老师1", "指导老师1电话", "指导老师2", "指导老师2电话",
    "乐团类别", "参展组别", "指定曲目", "自选曲目", "参展人数",
    *INSTRUMENTS,
    "合计", "备注", "用餐预约",
]


def seconds_to_human(seconds) -> str:
    """Match Laravel's ``SToHis`` helper (whole seconds, Chinese suffixes)."""

    try:
        value = int(seconds or 0)
    except (TypeError, ValueError):
        value = 0
    minutes, remainder = divmod(value, 60)
    return f"{minutes}分{remainder}秒" if minutes else f"{remainder}秒"


def status_label(status) -> str:
    return {0: "待审核", 1: "已通过", -1: "未通过"}.get(status, "-")


def group_number(group) -> Optional[int]:
    return {"小学组": 1, "中学组": 2, "大学组": 3}.get(group)


def export_code(group, report_id: int):
    """Build the report code used by ``Api\\ExportController``.

    Unknown group values are left as a string rather than producing an invalid
    integer.  The Laravel code's known production groups all map to 1..3.
    """

    number = group_number(group)
    if number is None:
        return "-"
    return int(f"240{number}000000") + int(report_id)


def _json_list(value):
    if value in (None, ""):
        return []
    if isinstance(value, (list, tuple)):
        return list(value)
    if isinstance(value, str):
        try:
            parsed = json.loads(value)
        except (TypeError, ValueError):
            return [value]
        return parsed if isinstance(parsed, list) else [parsed]
    return [value]


def _person_name(person):
    return getattr(person, "name", "") if person else ""


def _snapshot(reports: Iterable[Report]):
    """Materialise related objects once, preserving Laravel eager-load order."""

    records = list(reports)
    report_ids = [item.id for item in records]
    links = list(ReportPerson.objects.filter(report_id__in=report_ids).order_by("id")) if report_ids else []
    person_ids = [link.person_id for link in links if link.person_id is not None]
    people = {person.id: person for person in Person.objects.filter(id__in=person_ids)} if person_ids else {}
    user_ids = [item.user_id for item in records if item.user_id is not None]
    users = {user.id: user for user in User.all_objects.filter(id__in=user_ids)} if user_ids else {}
    file_ids = [x for item in records for x in (item.file, item.spectrum) if x is not None]
    files = {item.id: item for item in Files.objects.filter(id__in=file_ids)} if file_ids else {}
    grouped = defaultdict(list)
    for link in links:
        grouped[link.report_id].append((link, people.get(link.person_id)))
    return records, grouped, users, files


def _members(grouped, report_id):
    """Return names grouped by Laravel ``ReportPerson.position`` values."""

    result = defaultdict(list)
    for link, person in grouped.get(report_id, []):
        if person is not None:
            result[int(link.position)].append(person)
    return result


def _joined(members, position):
    return "、".join(_person_name(person) for person in members.get(position, []))


# --- 附件2 对齐的共用取值（与报名信息表 PDF 的 form_context 同口径） -----------

def _conductor(members):
    """指挥与指挥电话：全部指挥「、」相连（附件2 指挥栏 + 联系电话栏）。"""
    conductors = members.get(2, [])
    name = "、".join(_person_name(person) for person in conductors)
    phone = "、".join(str(getattr(p, "phone", "") or "") for p in conductors if getattr(p, "phone", ""))
    return name, phone


def _adviser_slots(teachers):
    """附件二固定两个指导老师名额槽：第 1 槽取第一位，其余（含超出 2 人）并入第 2 槽。"""
    if teachers:
        first = (_person_name(teachers[0]), str(getattr(teachers[0], "phone", "") or ""))
    else:
        first = ("", "")
    rest = teachers[1:]
    second = (
        "、".join(_person_name(p) for p in rest),
        "、".join(str(getattr(p, "phone", "") or "") for p in rest if getattr(p, "phone", "")),
    )
    return first, second


def _headcount(formal, reserve):
    """附件2「参展人数」栏的官方文案。"""
    return f"正式队员 {len(formal)} 人，预备队员 {len(reserve)} 人"


def _instrument_roster(members):
    """正式队员按乐器名单（附件2「正式队员名单」栏）：只列有人的乐器槽，换行分隔。"""
    names = {key: [] for key in INSTRUMENTS}
    for person in members.get(0, []):
        names[_instrument_bucket(getattr(person, "instrument", ""))].append(_person_name(person))
    return "\n".join(f"{key}：{'、'.join(values)}" for key, values in names.items() if values)


def _meal_and_remark(report):
    """用餐预约归位到官方 6 个时段，无法归位的条目并入备注（与报名信息表 PDF 同口径）。"""
    from apps.api.registration_form import MEALS, _meal_cells
    cells, leftovers = _meal_cells(report)
    meals = "、".join(label for label, mark in zip(MEALS, cells) if mark)
    parts = [str(report.remark or "")]
    if leftovers:
        parts.append("用餐预约：" + "、".join(leftovers))
    return meals, "\n".join(part for part in parts if part)


def report_data_rows(reports: Iterable[Report]) -> list[list]:
    """Rows for ``GET /export/data`` (the 19-column Laravel export)."""

    records, grouped, users, files = _snapshot(reports)
    rows = [REPORT_DATA_HEADINGS.copy()]
    for item in records:
        members = _members(grouped, item.id)
        user = users.get(item.user_id)
        spectrum = files.get(item.spectrum)
        video = files.get(item.file)
        rows.append([
            getattr(user, "nickname", "") if user else "",
            item.choir_name or "",
            item.name or "",
            item.name1 or "",
            export_code(item.group, item.id),
            item.school_name or "",
            item.desc or "",
            item.group or "",
            item.contact_name or "",
            item.contact_phone or "",
            item.contact_way or "",
            seconds_to_human(item.time_length),
            getattr(spectrum, "url", "") if spectrum else "",
            getattr(video, "url", "") if video else "",
            status_label(item.status),
            _joined(members, 0),
            _joined(members, 1),
            _joined(members, 2),
            _joined(members, 4),
        ])
    return rows


def admin_data1_rows(reports: Iterable[Report]) -> list[list]:
    """Rows for ``GET /admin/export/data1`` (报名数据，列集对齐附件2)."""

    records, grouped, users, _files = _snapshot(reports)
    rows = [ADMIN_DATA1_HEADINGS.copy()]
    for index, item in enumerate(records, start=1):
        members = _members(grouped, item.id)
        user = users.get(item.user_id)
        formal, reserve = members.get(0, []), members.get(1, [])
        conductor_name, conductor_phone = _conductor(members)
        (teacher1_name, teacher1_phone), (teacher2_name, teacher2_phone) = _adviser_slots(members.get(4, []))
        meals, remark = _meal_and_remark(item)
        rows.append([
            index,
            item.school_name or "",
            item.contact_name or "",
            item.contact_phone or "",
            conductor_name,
            conductor_phone,
            teacher1_name, teacher1_phone, teacher2_name, teacher2_phone,
            item.establishment or "",
            item.group or "",
            item.name1 or "",
            item.name or "",
            _headcount(formal, reserve),
            _instrument_roster(members),
            _joined(members, 1),
            remark,
            meals,
            getattr(user, "nickname", "") if user else "",
            item.choir_name or "",
            seconds_to_human(item.time_length),
            item.contact_way or "",
            item.desc or "",
            status_label(item.status),
        ])
    return rows


def _instrument_bucket(instrument):
    """Map known input spellings to the source export's instrument columns."""

    value = str(instrument or "").strip()
    aliases = {
        "低音单簧": "低音单簧管", "中音萨克": "中音萨克斯", "次中音萨": "次中音萨克斯",
        "上低音萨": "上低音萨克斯", "低音大提": "低音大提琴",
    }
    value = aliases.get(value, value)
    return value if value in INSTRUMENTS else "其他"


def admin_data2_rows(reports: Iterable[Report]) -> list[list]:
    """Rows for ``GET /admin/export/data2`` (器乐统计，列集对齐附件2).

    乐器 17 栏对应附件2「正式队员名单」的乐器槽，因此只统计正式队员
    （预备队员人数体现在「参展人数」栏）。合计即正式队员总数。
    """

    records, grouped, users, _files = _snapshot(reports)
    rows = [ADMIN_DATA2_HEADINGS.copy()]
    for index, item in enumerate(records, start=1):
        members = _members(grouped, item.id)
        user = users.get(item.user_id)
        instrument_counts = {key: 0 for key in INSTRUMENTS}
        for person in members.get(0, []):
            instrument_counts[_instrument_bucket(getattr(person, "instrument", ""))] += 1
        conductor_name, conductor_phone = _conductor(members)
        (teacher1_name, teacher1_phone), (teacher2_name, teacher2_phone) = _adviser_slots(members.get(4, []))
        meals, remark = _meal_and_remark(item)
        rows.append([
            index,
            getattr(user, "nickname", "") if user else "",
            item.school_name or "",
            item.choir_name or "",
            item.contact_name or "",
            item.contact_phone or "",
            conductor_name,
            conductor_phone,
            teacher1_name, teacher1_phone, teacher2_name, teacher2_phone,
            item.establishment or "",
            item.group or "",
            item.name1 or "",
            item.name or "",
            _headcount(members.get(0, []), members.get(1, [])),
            *(str(instrument_counts[key]) for key in INSTRUMENTS),
            sum(instrument_counts.values()),
            remark,
            meals,
        ])
    return rows


def _safe_title(title: str, fallback: str = "Sheet") -> str:
    title = str(title or fallback)
    # Excel worksheet names cannot contain these characters and are limited to
    # 31 characters.  Laravel's configured titles are already valid, but this
    # keeps old/custom data from making the response fail.
    for char in "[]:*?/\\":
        title = title.replace(char, "_")
    return title[:31] or fallback


def workbook_response(
    sheets: Sequence[tuple[str, Sequence[Sequence]]],
    filename: str,
    *,
    draw_style: bool = False,
) -> HttpResponse:
    """Serialize one or more row matrices into a downloadable XLSX response."""

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, Side
        from openpyxl.utils import get_column_letter

        workbook = Workbook()
        first = True
        for title, rows in sheets:
            sheet = workbook.active if first else workbook.create_sheet()
            first = False
            sheet.title = _safe_title(title)
            for row in rows:
                sheet.append(list(row))
            if draw_style:
                thin = Side(style="thin", color="000000")
                border = Border(left=thin, right=thin, top=thin, bottom=thin)
                max_row = max(1, sheet.max_row)
                for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=3):
                    for cell in row:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.border = border
                for row_index in range(1, max_row + 1):
                    sheet.row_dimensions[row_index].height = 25
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="000000")
                if max_row >= 2:
                    for cell in sheet[2]:
                        cell.font = Font(bold=True, color="000000")
                sheet.merge_cells("A1:C1")
                sheet.column_dimensions["A"].width = 10
                sheet.column_dimensions["B"].width = 50
                sheet.column_dimensions["C"].width = 10
            else:
                sheet.freeze_panes = "A2"
                for cell in sheet[1] if sheet.max_row else []:
                    cell.font = Font(bold=True)
                for column in sheet.columns:
                    values = [len(str(cell.value or "")) for cell in column]
                    width = min(50, max(12, max(values or [0]) + 2))
                    sheet.column_dimensions[get_column_letter(column[0].column)].width = width
        output = io.BytesIO()
        workbook.save(output)
        content = output.getvalue()
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    except ImportError:
        # Keep the same graceful degradation as the existing endpoint helper.
        import csv

        output = io.StringIO()
        writer = csv.writer(output)
        for _title, rows in sheets:
            writer.writerows(rows)
        content = output.getvalue().encode("utf-8-sig")
        content_type = "text/csv"
    result = HttpResponse(content, content_type=content_type)
    result["Content-Disposition"] = f'attachment; filename="{filename}"'
    return result


def reports_export_response(reports: Iterable[Report], filename: str = "数据导出.xlsx") -> HttpResponse:
    return workbook_response([("数据导出", report_data_rows(reports))], filename)


def admin_data1_response(reports: Iterable[Report], filename: str = "报名数据.xlsx") -> HttpResponse:
    return workbook_response([("报名数据", admin_data1_rows(reports))], filename)


def admin_data2_response(reports: Iterable[Report], filename: str = "器乐统计数据.xlsx") -> HttpResponse:
    return workbook_response([("器乐统计数据", admin_data2_rows(reports))], filename)


DRAW_TYPES = {
    1: "大学组 (非专业组)",
    2: "大学组（专业组）",
    3: "教师组",
    4: "中小学组",
}


def draw_type_label(draw_type) -> str:
    try:
        return DRAW_TYPES.get(int(draw_type), "")
    except (TypeError, ValueError):
        return ""


def draw_rows(draws: Iterable[Draw], draw_type=None) -> list[list]:
    rows = [["序号", "合唱团名称", "抽签序号"]]
    for index, item in enumerate(draws, start=1):
        rows.append([index, item.name or "", item.index if item.index is not None else ""])
    return rows


def draw_response(draws: Iterable[Draw], draw_type: int, filename: Optional[str] = None) -> HttpResponse:
    label = draw_type_label(draw_type)
    rows = [[f"{label}现场展演抽签顺序表", "", ""]] + draw_rows(draws, draw_type)
    return workbook_response([(f"{label}现场展演抽签顺序表", rows)], filename or f"{label}现场展演抽签顺序表.xlsx", draw_style=True)


def draw_all_response(draws_by_type: Mapping[int, Iterable[Draw]], filename: str = "所有类别抽签排序表.xlsx") -> HttpResponse:
    sheets = []
    titles = {1: "大学组 (非专业组)", 2: "大学组（专业组）", 3: "教师组", 4: "中小学组"}
    for draw_type in (1, 2, 3, 4):
        label = draw_type_label(draw_type)
        rows = [[f"{label}现场展演抽签顺序表", "", ""]] + draw_rows(draws_by_type.get(draw_type, []), draw_type)
        sheets.append((titles[draw_type], rows))
    return workbook_response(sheets, filename, draw_style=True)

