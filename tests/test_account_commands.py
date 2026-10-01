"""账号批量导入（import_accounts）与后台建号（create_staff_accounts）的命令契约。"""
import tempfile
from io import StringIO
from unittest.mock import patch

from django.core.management import call_command
from django.core.management.base import CommandError
from openpyxl import Workbook, load_workbook

from apps.core.models import User

from .base import ApiTestCase

TEMPLATE_HINT_ROW = ["账号（登录用，唯一，不可与现有账号重复）",
                     "账号名称（学校/单位显示名）",
                     "所属市州（填市州登录账号，如 city01；中小学账号一般要填，市州账号自己不用填）",
                     "账号类型（填数字）：0=高校，1=市州，2=组委会，3=管理员，5=中小学",
                     "合并办学特许（仅中小学=5 可填）：可同时报小学组和中学组填 1，不可以就留空",
                     "初始密码（留空用默认 scylb@2026；不同账号要不同密码就逐行填）"]


def write_excel(path, rows, hint_row=True):
    """按模板结构写一个待导入的 Excel：第一行提示行（可选），随后是数据行。"""
    workbook = Workbook()
    sheet = workbook.active
    if hint_row:
        sheet.append(TEMPLATE_HINT_ROW)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


class ImportAccountsMixin(ApiTestCase):
    def run_import(self, rows, *args, hint_row=True):
        out = StringIO()
        with tempfile.TemporaryDirectory() as tmp:
            path = write_excel("%s/导入.xlsx" % tmp, rows, hint_row=hint_row)
            call_command("import_accounts", path, stdout=out, *args)
        return out.getvalue()

    def run_import_expect_error(self, rows, *args, hint_row=True):
        out = StringIO()
        with tempfile.TemporaryDirectory() as tmp:
            path = write_excel("%s/导入.xlsx" % tmp, rows, hint_row=hint_row)
            with self.assertRaises(CommandError) as ctx:
                call_command("import_accounts", path, stdout=out, *args)
        return out.getvalue(), str(ctx.exception)


class ImportAccountsTemplateTests(ImportAccountsMixin):
    def test_template_first_row_hints_type_values(self):
        out = StringIO()
        with tempfile.TemporaryDirectory() as tmp:
            path = "%s/模板.xlsx" % tmp
            call_command("import_accounts", "--template", path, stdout=out)
            sheet = load_workbook(path).active

        self.assertIn("账号", sheet["A1"].value)
        self.assertIn("账号名称", sheet["B1"].value)
        # C1：所属市州填登录账号，且提示市州自己不用填
        self.assertIn("所属市州", sheet["C1"].value)
        self.assertIn("登录账号", sheet["C1"].value)
        self.assertIn("市州账号自己不用填", sheet["C1"].value)
        # D1：账号类型填法（用户方约定的 5 种）
        self.assertIn("账号类型", sheet["D1"].value)
        for expected in ("0=高校", "1=市州", "2=组委会", "3=管理员", "5=中小学"):
            self.assertIn(expected, sheet["D1"].value)
        # E1 说明合并办学特许怎么填：1=可以，不可以留空
        self.assertIn("合并办学", sheet["E1"].value)
        self.assertIn("填 1", sheet["E1"].value)
        self.assertIn("留空", sheet["E1"].value)
        # F1 说明初始密码怎么填：留空用默认
        self.assertIn("初始密码", sheet["F1"].value)
        self.assertIn("留空", sheet["F1"].value)
        self.assertIn("scylb@2026", sheet["F1"].value)
        self.assertIn("模板已生成", out.getvalue())


class ImportAccountsCommandTests(ImportAccountsMixin):
    def test_import_creates_accounts_with_default_password(self):
        output = self.run_import([
            ["school-a", "某某中学", None, 5],
            ["uni-a", "某某大学", None, "0"],          # 类型传文本数字也要认
            ["committee-a", "组委会组", None, 2],
        ])

        school = User.all_objects.get(username="school-a")
        self.assertEqual((school.type, school.nickname), (5, "某某中学"))
        self.assertTrue(school.check_password("scylb@2026"))
        self.assertEqual(User.all_objects.get(username="uni-a").type, 0)
        self.assertEqual(User.all_objects.get(username="committee-a").type, 2)
        self.assertIn("新建 3 个", output)

    def test_numeric_username_cell_stored_without_float_tail(self):
        # Excel 把纯数字账号存成浮点：不能落成 "123456.0"
        self.run_import([[123456, "数字账号学校", None, 5]])

        self.assertTrue(User.all_objects.filter(username="123456", type=5).exists())

    def test_existing_username_is_skipped_not_failed(self):
        self.create_user("dup-school", 0)

        output = self.run_import([
            ["dup-school", "重复的学校", None, 0],
            ["fresh-school", "新学校", None, 5],
        ])

        self.assertEqual(User.all_objects.filter(type=5, username="fresh-school").count(), 1)
        self.assertIn("跳过（已存在）1 个：dup-school", output)
        self.assertIn("新建 1 个", output)
        # 没有失败行：不抛 CommandError（上面 run_import 未捕获即为证）

    def test_in_file_duplicate_fails_second_occurrence(self):
        output, error = self.run_import_expect_error([
            ["twice-a", "第一个", None, 0],
            ["twice-a", "第二个", None, 0],
        ])

        self.assertEqual(User.all_objects.filter(username="twice-a").count(), 1)
        self.assertEqual(User.all_objects.get(username="twice-a").nickname, "第一个")
        self.assertIn("失败 1 行", output)
        self.assertIn("第 3 行", output)
        self.assertIn("重复", output)
        self.assertIn("有 1 行未入库", error)

    def test_invalid_type_row_fails_but_valid_rows_still_import(self):
        output, error = self.run_import_expect_error([
            ["good-a", "好学校", None, 5],
            ["bad-type", "类型错", None, 4],           # 省级是历史渠道，不开放导入
            ["bad-text", "非数字", None, "abc"],
        ])

        self.assertTrue(User.all_objects.filter(username="good-a", type=5).exists())
        self.assertFalse(User.all_objects.filter(username="bad-type").exists())
        # 行级明细在 stdout，异常里是摘要
        self.assertIn("失败 2 行", output)
        self.assertIn("不在可导入范围内", output)
        self.assertIn("必须是数字", output)
        self.assertIn("有 2 行未入库", error)

    def test_missing_fields_report_row_number(self):
        output, error = self.run_import_expect_error([
            ["", "没有账号", None, 0],
            ["no-nick", "", None, 0],
            ["x" * 31, "超长账号", None, 0],
            ["带 空格", "空格账号", None, 0],
        ])

        self.assertFalse(User.all_objects.exists())
        for row_no in (2, 3, 4, 5):
            self.assertIn("第 %d 行" % row_no, output)
        self.assertIn("账号不能为空", output)
        self.assertIn("账号名称不能为空", output)
        self.assertIn("超过 30 个字符", output)
        self.assertIn("不能含空格", output)

    def test_empty_sheet_rejected(self):
        output, error = self.run_import_expect_error([], hint_row=True)
        self.assertIn("没有可导入的数据行", error)

    def test_parent_option_links_to_city_account(self):
        city = self.create_user("city-imp", 1)

        output = self.run_import([["primary-a", "挂靠学校", None, 5]], "--parent", "city-imp")

        school = User.all_objects.get(username="primary-a")
        self.assertEqual(school.parent_id, city.id)
        self.assertIn("挂靠市州 1 个", output)
        self.assertIn("city-imp：primary-a", output)

    def test_parent_option_rejects_missing_account_before_writing(self):
        output, error = self.run_import_expect_error(
            [["primary-b", "没挂上的学校", None, 5]], "--parent", "ghost-city")

        self.assertFalse(User.all_objects.filter(username="primary-b").exists())
        self.assertIn("--parent 指定的账号不存在", error)

    def test_custom_password_option(self):
        self.run_import([["pw-a", "自定义密码学校", None, 5]], "--password", "另@2026")

        self.assertTrue(User.all_objects.get(username="pw-a").check_password("另@2026"))

    def test_can_report_twice_column(self):
        self.run_import([
            ["merge-a", "九年一贯制学校", None, 5, 1],
            ["float-a", "浮点特许学校", None, 5, 1.0],   # Excel 常把数字存成浮点，不能落成 "1.0"
            ["plain-a", "普通学校", None, 5, None],      # 留空 = 不可以
            ["zero-a", "填零学校", None, 5, 0],          # 0 视同留空
        ])

        self.assertTrue(User.all_objects.get(username="merge-a").can_report_twice)
        self.assertTrue(User.all_objects.get(username="float-a").can_report_twice)
        self.assertFalse(User.all_objects.get(username="plain-a").can_report_twice)
        self.assertFalse(User.all_objects.get(username="zero-a").can_report_twice)

    def test_can_report_twice_only_for_primary_secondary(self):
        # 特许只发给中小学：其他类型填 1 该行报错，同文件的合法行照常入库
        output, error = self.run_import_expect_error([
            ["uni-twice", "高校也想报两组", None, 0, 1],
            ["primary-ok", "中小学正常特许", None, 5, 1],
        ])

        self.assertFalse(User.all_objects.filter(username="uni-twice").exists())
        self.assertTrue(User.all_objects.get(username="primary-ok").can_report_twice)
        self.assertIn("只有中小学账号", output)
        self.assertIn("有 1 行未入库", error)

    def test_can_report_twice_rejects_ambiguous_values(self):
        output, error = self.run_import_expect_error([
            ["bad-two", "填2学校", None, 5, 2],
            ["bad-text", "填字学校", None, 5, "可以"],
        ])

        self.assertEqual(
            User.all_objects.filter(username__in=["bad-two", "bad-text"]).count(), 0)
        self.assertIn("只能填 1 或留空", output)
        self.assertIn("有 2 行未入库", error)

    def test_city_column_links_parent_and_falls_back_to_option(self):
        city = self.create_user("city-col", 1)

        output = self.run_import([
            ["p-city", "挂列学校", "city-col", 5, None],    # 列里填了：以列为准
            ["p-opt", "靠参数学校", None, 5, None],          # 列空：回退 --parent
        ], "--parent", "city-col")

        self.assertEqual(User.all_objects.get(username="p-city").parent_id, city.id)
        self.assertEqual(User.all_objects.get(username="p-opt").parent_id, city.id)
        self.assertIn("挂靠市州 1 个", output)
        self.assertIn("city-col：p-city、p-opt", output)

    def test_city_column_wins_over_parent_option(self):
        self.create_user("city-fb", 1)
        chosen = self.create_user("city-ch", 1)

        self.run_import([["p-both", "两处都填的学校", "city-ch", 5, None]],
                        "--parent", "city-fb")

        self.assertEqual(User.all_objects.get(username="p-both").parent_id, chosen.id)

    def test_city_column_missing_or_non_city_fails_row(self):
        self.create_user("school-x", 0)

        output, error = self.run_import_expect_error([
            ["p-ghost", "市州写错", "ghost-city", 5, None],
            ["p-notcity", "指到高校", "school-x", 5, None],
            ["p-fine", "正常学校", None, 5, None],
        ])

        self.assertFalse(
            User.all_objects.filter(username__in=["p-ghost", "p-notcity"]).exists())
        self.assertTrue(User.all_objects.filter(username="p-fine").exists())
        self.assertIn("所属市州账号不存在：ghost-city", output)
        self.assertIn("所属市州账号不是市州（1）：school-x", output)
        self.assertIn("有 2 行未入库", error)

    def test_city_type_row_must_not_fill_city_column(self):
        output, error = self.run_import_expect_error(
            [["city-new", "新市州", "city-imp", 1, None]])

        self.assertFalse(User.all_objects.filter(username="city-new").exists())
        self.assertIn("市州账号（1）不需要填所属市州", output)
        self.assertIn("有 1 行未入库", error)

    def test_password_column_overrides_default_per_row(self):
        output = self.run_import([
            ["pw-x", "自定义密码学校", None, 5, None, "特@2026"],
            ["pw-y", "默认密码学校", None, 5, None, None],      # F 列留空 → 用 --password/默认
        ])

        self.assertTrue(User.all_objects.get(username="pw-x").check_password("特@2026"))
        self.assertTrue(User.all_objects.get(username="pw-y").check_password("scylb@2026"))
        self.assertIn("初始密码分布", output)
        self.assertIn("scylb@2026 ×1 个", output)
        self.assertIn("特@2026 ×1 个", output)

    def test_city_reference_to_row_in_same_file(self):
        # C 列可以引用同一张表里的市州行：命令先建市州、再建挂靠它的学校
        output = self.run_import([
            ["p-inline", "挂表内市州的学校", "city-inline", 5, None, None],
            ["city-inline", "表内市州", None, 1, None, None],   # 市州行故意放在学校后面
        ])

        city = User.all_objects.get(username="city-inline")
        self.assertEqual(User.all_objects.get(username="p-inline").parent_id, city.id)
        self.assertIn("city-inline：p-inline", output)

    def test_city_reference_to_non_city_row_in_same_file(self):
        output, error = self.run_import_expect_error([
            ["uni-inline", "表内高校", None, 0, None, None],
            ["p-bad", "挂到了高校行", "uni-inline", 5, None, None],
        ])

        self.assertFalse(User.all_objects.filter(username="p-bad").exists())
        self.assertTrue(User.all_objects.filter(username="uni-inline").exists())
        self.assertIn("所属市州账号不是市州（1）：uni-inline", output)
        self.assertIn("有 1 行未入库", error)


class CreateStaffAccountsTests(ApiTestCase):
    def run_create(self, *args):
        out = StringIO()
        call_command("create_staff_accounts", stdout=out, *args)
        return out.getvalue()

    def test_default_creates_three_committee_and_three_admin(self):
        output = self.run_create()

        self.assertEqual(User.all_objects.filter(type=2).count(), 3)
        self.assertEqual(User.all_objects.filter(type=3).count(), 3)
        committee = User.all_objects.filter(type=2).order_by("username")
        self.assertEqual(list(committee.values_list("username", flat=True)),
                         ["committee01", "committee02", "committee03"])
        self.assertEqual(list(committee.values_list("nickname", flat=True)),
                         ["组委会01", "组委会02", "组委会03"])
        admin = User.all_objects.get(username="admin01")
        self.assertEqual(admin.nickname, "管理员01")
        self.assertTrue(admin.check_password("scylb@2026"))
        self.assertIn("共创建 6 个账号", output)

    def test_existing_numbers_are_skipped(self):
        # 复刻共享库现状：committee01/admin01 已存在
        self.create_user("committee01", 2)
        self.create_user("admin01", 3)

        self.run_create()

        self.assertEqual(
            set(User.all_objects.filter(type=2).values_list("username", flat=True)),
            {"committee01", "committee02", "committee03", "committee04"})
        self.assertEqual(
            set(User.all_objects.filter(type=3).values_list("username", flat=True)),
            {"admin01", "admin02", "admin03", "admin04"})

    def test_rerun_creates_next_batch_without_collision(self):
        self.run_create()
        output = self.run_create()

        self.assertEqual(
            set(User.all_objects.filter(type=2).values_list("username", flat=True)),
            {"committee01", "committee02", "committee03",
             "committee04", "committee05", "committee06"})
        self.assertIn("共创建 6 个账号", output)

    def test_soft_deleted_account_still_occupies_its_number(self):
        # username 唯一约束不分死活：软删的 committee01 也占号，不许复用
        self.create_user("committee01", 2)
        User.all_objects.get(username="committee01").delete()

        self.run_create()

        numbers = set(User.all_objects.filter(
            username__startswith="committee").values_list("username", flat=True))
        self.assertEqual(numbers, {"committee01", "committee02", "committee03", "committee04"})

    def test_type_and_count_options(self):
        output = self.run_create("--type", "2", "--count", "1", "--password", "另@2026")

        self.assertEqual(User.all_objects.filter(type=2).count(), 1)
        self.assertEqual(User.all_objects.filter(type=3).count(), 0)
        self.assertTrue(User.all_objects.get(username="committee01").check_password("另@2026"))
        self.assertIn("共创建 1 个账号", output)
