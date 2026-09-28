import io
import unittest
from datetime import timedelta

from django.test import override_settings
from django.utils import timezone

from apps.core.models import PersonalAccessToken, User

from .base import ApiTestCase

try:
    import bcrypt  # noqa: F401
    HAVE_BCRYPT = True
except ImportError:
    HAVE_BCRYPT = False

# A real Laravel-style bcrypt digest of the password "pw", generated with:
#   python -c "import bcrypt; h=bcrypt.hashpw(b'pw', bcrypt.gensalt(10, prefix=b'2b')).decode(); print('$2y$'+h[4:])"
LARAVEL_PW_HASH = "$2y$10$m0OVXEJF5/UUuGOrXFIkgOJDrecupQgpaIpU.V59SMl3iLseuwhKy"


class AuthenticationContractTests(ApiTestCase):
    def setUp(self):
        self.school = self.create_user("school", 0, nickname="学校")

    def test_login_rejects_bad_password_without_issuing_token(self):
        result = self.json_request(
            "post", "/api/login", {"username": "school", "password": "wrong"}
        )

        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json(), {"code": 1, "msg": "账号或密码错误", "data": None})
        self.assertFalse(PersonalAccessToken.objects.exists())

    def test_missing_and_invalid_bearer_tokens_are_unauthorized(self):
        missing = self.client.get("/api/file/list")
        invalid = self.client.get(
            "/api/file/list", HTTP_AUTHORIZATION="Bearer definitely-invalid"
        )

        self.assertEqual(missing.status_code, 401)
        self.assertEqual(invalid.status_code, 401)

    def test_logout_revokes_only_the_presented_token(self):
        first, first_plain = PersonalAccessToken.issue(self.school)
        second, second_plain = PersonalAccessToken.issue(self.school)
        self.client.defaults["HTTP_AUTHORIZATION"] = "Bearer " + first_plain

        result = self.client.post("/api/logout")

        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()["code"], 0)
        self.assertFalse(PersonalAccessToken.objects.filter(pk=first.pk).exists())
        self.assertTrue(PersonalAccessToken.objects.filter(pk=second.pk).exists())
        self.assertEqual(self.client.get("/api/file/list").status_code, 401)
        self.client.defaults["HTTP_AUTHORIZATION"] = "Bearer " + second_plain
        self.assertEqual(self.client.get("/api/file/list").status_code, 200)

    @override_settings(TOKEN_TTL_HOURS=4)
    def test_expired_token_is_rejected_and_removed(self):
        token, plain = PersonalAccessToken.issue(self.school)
        PersonalAccessToken.objects.filter(pk=token.pk).update(
            last_used_at=timezone.now() - timedelta(hours=5)
        )
        self.client.defaults["HTTP_AUTHORIZATION"] = "Bearer " + plain

        result = self.client.get("/api/file/list")

        self.assertEqual(result.status_code, 401)
        self.assertFalse(PersonalAccessToken.objects.filter(pk=token.pk).exists())


class RolePermissionTests(ApiTestCase):
    def setUp(self):
        self.users = {
            0: self.create_user("school", 0),
            1: self.create_user("city", 1),
            2: self.create_user("committee", 2),
            3: self.create_user("admin", 3),
            4: self.create_user("province", 4),
            5: self.create_user("primary", 5),
        }

    def test_each_role_can_reach_its_own_dashboard_only(self):
        routes = {
            0: "/api/school/index/total",
            1: "/api/city/index/total",
            2: "/api/committee/index/total",
            3: "/api/admin/index/total",
            4: "/api/province/index/total",
            5: "/api/primary/index/total",
        }

        for user_type, own_route in routes.items():
            with self.subTest(user_type=user_type, route=own_route):
                self.authorize_as(self.users[user_type])
                own = self.client.get(own_route)
                self.assertEqual(own.status_code, 200)
                self.assertEqual(own.json()["code"], 0)
                other_route = routes[(user_type + 1) % len(routes)]
                denied = self.client.get(other_route)
                self.assertEqual(denied.status_code, 403)
                self.assertEqual(denied.json()["code"], "FORBIDDEN")
                self.assertEqual(denied.json()["msg"], "无该页面操作权限！")

    def test_percent_dashboards_apply_the_same_role_middleware(self):
        self.authorize_as(self.users[1])

        school = self.client.get("/api/school/index/percent")
        province = self.client.get("/api/province/index/percent")

        self.assertEqual(school.status_code, 403)
        self.assertEqual(province.status_code, 403)

    def test_profile_update_cannot_modify_another_account(self):
        attacker = self.users[0]
        victim = self.users[1]
        self.authorize_as(attacker)

        result = self.json_request(
            "put", "/api/user", {"id": victim.id, "nickname": "被篡改"}
        )

        victim.refresh_from_db()
        self.assertEqual(victim.nickname, "city")
        if result.status_code == 200:
            self.assertEqual(result.json()["code"], 1)
        else:
            self.assertEqual(result.status_code, 403)
            self.assertEqual(result.json()["code"], "FORBIDDEN")
            self.assertEqual(result.json()["msg"], "无该页面操作权限！")

    def test_profile_update_cannot_escalate_own_role(self):
        user = self.users[0]
        self.authorize_as(user)

        result = self.json_request(
            "put", "/api/user", {"id": user.id, "nickname": "新名称", "type": 3}
        )

        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json()["code"], 0)
        user.refresh_from_db()
        self.assertEqual(user.nickname, "新名称")
        self.assertEqual(user.type, 0)

    def test_admin_and_committee_report_extensions_are_role_gated(self):
        report = self.make_report(self.users[0])
        routes = [
            f"/api/admin/report/{report.id}",
            "/api/admin/report/update",
            f"/api/committee/report/{report.id}",
            "/api/committee/report/update",
        ]

        for route in routes:
            expected = 3 if route.startswith("/api/admin/") else 2
            for user_type in range(6):
                if user_type == expected:
                    continue
                with self.subTest(route=route, user_type=user_type):
                    self.authorize_as(self.users[user_type])
                    if route.endswith("update"):
                        denied = self.json_request("put", route, {"id": report.id})
                    else:
                        denied = self.client.get(route)
                    self.assertEqual(denied.status_code, 403)
                    self.assertEqual(denied.json()["code"], "FORBIDDEN")
                    self.assertEqual(denied.json()["msg"], "无该页面操作权限！")


class PrimaryAndCityScopeTests(ApiTestCase):
    """中小学端（type=5）承接原市州端报名功能；市州端（type=1）降级为纯只读。"""

    def setUp(self):
        self.primary = self.create_user("primary", 5)
        self.city = self.create_user("city", 1)

    def test_primary_dashboard_mirrors_city_panels(self):
        report = self.make_report(self.primary, group="小学组")
        self.authorize_as(self.primary)
        for route in ("/api/primary/index/total", "/api/primary/index/percent",
                      "/api/primary/index/establishment"):
            result = self.client.get(route)
            self.assertEqual(result.status_code, 200)
            self.assertEqual(result.json()["code"], 0)
        # 与市州端同一口径：只统计本账号报送的报名
        rows = self.client.get("/api/primary/index/establishment").json()["data"]["data"]
        band = next(row for row in rows if row["name"] == "管乐团")
        self.assertEqual(band["total"], 1)   # make_report 默认管乐团

    def test_primary_endpoints_reject_other_roles(self):
        self.authorize_as(self.city)
        for route in ("/api/primary/index/total", "/api/primary/index/percent",
                      "/api/primary/index/establishment", "/api/primary/report/list",
                      "/api/primary/recommend/list"):
            result = self.client.get(route)
            self.assertEqual(result.status_code, 403, route)
            self.assertEqual(result.json()["code"], "FORBIDDEN")

    def test_city_read_only_endpoints_still_serve_type_one(self):
        report = self.make_report(self.city, group="小学组")
        self.authorize_as(self.city)
        self.assertEqual(self.client.get("/api/city/report/list").status_code, 200)
        self.assertEqual(self.client.get("/api/city/report/%s" % report.id).status_code, 200)
        self.assertEqual(self.client.get("/api/city/recommend/list").status_code, 200)

    def test_primary_scope_is_own_reports_only(self):
        other = self.create_user("primary2", 5)
        self.make_report(other, group="小学组")
        mine = self.make_report(self.primary, group="中学组")
        self.authorize_as(self.primary)

        payload = self.client.get("/api/primary/report/list").json()

        ids = [int(item["id"]) for item in payload["data"]]
        self.assertIn(mine.id, ids)
        self.assertNotIn(other.id, ids)


class AdminUserManagementPrimaryTests(ApiTestCase):
    """管理员账号管理覆盖中小学端（type=5）：列表、创建、账号导出。"""

    def setUp(self):
        self.admin = self.create_user("admin", 3)

    def test_admin_user_list_includes_primary_accounts(self):
        school = self.create_user("school", 0)
        city = self.create_user("city", 1)
        primary = self.create_user("primary", 5)
        committee = self.create_user("committee", 2)
        self.authorize_as(self.admin)

        ids = [item["id"] for item in self.client.get("/api/admin/user/list").json()["data"]]

        # 组委会账号自 2026-09-28 起也在管理员侧列表可见
        for user in (school, city, primary, committee):
            self.assertIn(user.id, ids)
        self.assertNotIn(self.admin.id, ids)

    def test_admin_creates_primary_account_that_can_access_primary_routes(self):
        self.authorize_as(self.admin)
        created = self.json_request("post", "/api/admin/user/", {
            "username": "primary-school", "password": "pw", "nickname": "某小学", "type": 5,
        })
        self.assertEqual(created.json()["code"], 0, created.content)
        self.assertEqual(User.objects.get(username="primary-school").type, 5)

        login = self.json_request(
            "post", "/api/login", {"username": "primary-school", "password": "pw"})
        self.client.defaults["HTTP_AUTHORIZATION"] = "Bearer " + login.json()["data"]["token"]
        self.assertEqual(self.client.get("/api/primary/index/total").json()["code"], 0)
        # 校级（高校端）路由对中小学端账号不可见
        self.assertEqual(self.client.get("/api/school/report/list").status_code, 403)

    def test_admin_account_export_includes_city_and_committee_accounts(self):
        self.create_user("school", 0)
        self.create_user("city", 1)
        self.create_user("primary", 5)
        self.create_user("committee", 2)
        self.authorize_as(self.admin)

        result = self.client.get("/api/admin/user/export")

        self.assertEqual(result.status_code, 200)
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl 未安装，无法解析导出文件")
        rows = list(load_workbook(io.BytesIO(result.content)).active.iter_rows(values_only=True))
        usernames = {row[0] for row in rows[1:]}
        # 导出与列表同口径（2026-09-28 起）：市州、组委会都进管理员侧导出
        for name in ("school", "city", "primary", "committee"):
            self.assertIn(name, usernames)
        self.assertNotIn("admin", usernames)


class CommitteeAccountVisibilityTests(ApiTestCase):
    """组委会账号（type=2）只在管理员侧的用户列表可见（2026-09-28 起）。

    此前 allowed_types 固定 (0,1,5)，管理员创建组委会账号后列表里找不到它，
    无从重置密码/修改；修复后管理员侧放开，组委会侧维持不见其他组委会账号。
    """

    def setUp(self):
        self.admin = self.create_user("admin", 3)
        self.committee_caller = self.create_user("committee-caller", 2)

    def test_admin_list_shows_committee_account_and_search_finds_it(self):
        committee = self.create_user("new-committee", 2, nickname="省组委会")
        self.authorize_as(self.admin)

        ids = [item["id"] for item in self.client.get("/api/admin/user/list").json()["data"]]
        self.assertIn(committee.id, ids)

        found = self.client.get("/api/admin/user/list", {"keyword": "new-committee"}).json()["data"]
        self.assertEqual([item["id"] for item in found], [committee.id])

    def test_admin_type_filter_can_narrow_to_committee_accounts(self):
        school = self.create_user("school", 0)
        committee = self.create_user("new-committee", 2)
        self.authorize_as(self.admin)

        data = self.client.get("/api/admin/user/list", {"type": 2}).json()["data"]
        ids = [item["id"] for item in data]

        # type=2 精确收窄：新建的组委会账号和 setUp 里的组委会调用者都在列，学校不在
        self.assertIn(committee.id, ids)
        self.assertIn(self.committee_caller.id, ids)
        self.assertNotIn(school.id, ids)
        self.assertEqual({item["type"] for item in data}, {2})

    def test_committee_list_still_hides_committee_accounts(self):
        school = self.create_user("school", 0)
        committee = self.create_user("new-committee", 2)
        self.authorize_as(self.committee_caller)

        ids = [item["id"] for item in self.client.get("/api/committee/user/list").json()["data"]]

        self.assertIn(school.id, ids)
        self.assertNotIn(committee.id, ids)

    def test_committee_type_filter_cannot_surface_committee_accounts(self):
        self.create_user("new-committee", 2)
        self.authorize_as(self.committee_caller)

        ids = [item["id"] for item in self.client.get("/api/committee/user/list", {"type": 2}).json()["data"]]

        self.assertEqual(ids, [])

    def test_committee_export_excludes_admin_and_committee_accounts(self):
        """组委会导出与组委会列表同口径：只有 0/1/5（修复前是 User.objects.all() 整表导出）。"""
        self.create_user("school", 0)
        self.create_user("city", 1)
        self.create_user("primary", 5)
        self.create_user("new-committee", 2)
        self.authorize_as(self.committee_caller)

        result = self.client.get("/api/committee/user/export")

        self.assertEqual(result.status_code, 200)
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl 未安装，无法解析导出文件")
        rows = list(load_workbook(io.BytesIO(result.content)).active.iter_rows(values_only=True))
        usernames = {row[0] for row in rows[1:]}
        for name in ("school", "city", "primary"):
            self.assertIn(name, usernames)
        for name in ("new-committee", "committee-caller", "admin"):
            self.assertNotIn(name, usernames)

    def test_committee_update_cannot_modify_admin_account(self):
        """组委会改号请求打到管理员头上：与「用户不存在」同判，昵称/类型/密码原封不动。"""
        admin = self.create_user("admin-victim", 3)
        original_nickname, original_password, original_type = admin.nickname, admin.password, admin.type
        self.authorize_as(self.committee_caller)

        result = self.json_request(
            "put", "/api/committee/user/",
            {"id": admin.id, "nickname": "被篡改", "password": "x", "type": 0},
        )

        self.assertEqual(result.json()["code"], 1)
        admin.refresh_from_db()
        self.assertEqual(admin.nickname, original_nickname)
        self.assertEqual(admin.type, original_type)
        self.assertEqual(admin.password, original_password)

        # 对照：组委会对自己范围内的学校账号仍可正常修改
        school = self.create_user("school", 0)
        update = self.json_request("put", "/api/committee/user/", {"id": school.id, "nickname": "改好了"})
        self.assertEqual(update.json()["code"], 0)
        school.refresh_from_db()
        self.assertEqual(school.nickname, "改好了")

    def test_committee_delete_cannot_delete_admin_account(self):
        admin = self.create_user("admin-victim", 3)
        school = self.create_user("school", 0)
        self.authorize_as(self.committee_caller)

        result = self.json_request("delete", "/api/committee/user/", {"ids": [admin.id, school.id]})

        self.assertEqual(result.json()["code"], 0)
        school.refresh_from_db()
        self.assertIsNotNone(school.deleted_at)   # 范围内的学校账号正常删
        admin.refresh_from_db()
        self.assertIsNone(admin.deleted_at)        # 管理员账号毫发无损

    def test_committee_info_cannot_read_admin_account(self):
        admin = self.create_user("admin-victim", 3)
        school = self.create_user("school", 0)
        self.authorize_as(self.committee_caller)

        self.assertIsNone(self.client.get("/api/committee/user/%s" % admin.id).json()["data"])
        found = self.client.get("/api/committee/user/%s" % school.id).json()
        self.assertEqual(found["data"]["id"], school.id)


@unittest.skipUnless(HAVE_BCRYPT, "bcrypt 未安装：pip install -r requirements.txt")
class LegacyBcryptLoginTests(ApiTestCase):
    """存量 Laravel bcrypt 用户登录：兼容校验 + 成功后透明升级 PBKDF2。"""

    def _legacy_user(self):
        # Bypass create_user()/set_password() to mirror import_laravel_data.py,
        # which stores the raw bcrypt digest exactly as found in the old database.
        user = User(username="legacy", nickname="老系统", type=0)
        user.password = LARAVEL_PW_HASH
        user.save()
        return user

    def test_login_with_laravel_hash_succeeds_and_upgrades_to_pbkdf2(self):
        user = self._legacy_user()
        before_updated_at = user.updated_at

        result = self.json_request("post", "/api/login", {"username": "legacy", "password": "pw"})

        self.assertEqual(result.status_code, 200)
        payload = result.json()
        self.assertEqual(payload["code"], 0)
        self.assertEqual(payload["msg"], "登录成功")
        self.assertRegex(payload["data"]["token"], r"^\d+\|[0-9a-f]{40}$")
        self.assertEqual(payload["data"]["user"], {
            "id": user.id, "username": "legacy", "nickname": "老系统",
            "description": "", "tel": "", "leader": "", "type": 0,
            "can_report_twice": False, "parent_id": None,
        })
        self.assertTrue(PersonalAccessToken.objects.filter(tokenable_id=user.id).exists())
        user.refresh_from_db()
        self.assertTrue(user.password.startswith("pbkdf2_sha256$"))
        # A hash reformat is not a content change: updated_at must not move.
        self.assertEqual(user.updated_at, before_updated_at)

    def test_wrong_password_is_rejected_and_hash_untouched(self):
        user = self._legacy_user()

        result = self.json_request("post", "/api/login", {"username": "legacy", "password": "wrong"})

        self.assertEqual(result.status_code, 200)
        self.assertEqual(result.json(), {"code": 1, "msg": "账号或密码错误", "data": None})
        self.assertFalse(PersonalAccessToken.objects.exists())
        user.refresh_from_db()
        self.assertEqual(user.password, LARAVEL_PW_HASH)

    def test_upgraded_user_and_native_pbkdf2_user_relogin(self):
        legacy = self._legacy_user()
        self.json_request("post", "/api/login", {"username": "legacy", "password": "pw"})
        legacy.refresh_from_db()
        upgraded_hash = legacy.password

        second = self.json_request("post", "/api/login", {"username": "legacy", "password": "pw"})
        self.assertEqual(second.json()["code"], 0)
        legacy.refresh_from_db()
        # Second login goes through the plain Django path and must not rewrite.
        self.assertEqual(legacy.password, upgraded_hash)

        modern = self.create_user("modern", 0)
        result = self.json_request("post", "/api/login", {"username": "modern", "password": "pw"})
        self.assertEqual(result.json()["code"], 0)
        modern.refresh_from_db()
        self.assertTrue(modern.password.startswith("pbkdf2_sha256$"))


class CurrentUserInfoTests(ApiTestCase):
    """GET /api/user：只认 token，返回当前登录用户自己（复用 user_dict）。"""

    def setUp(self):
        self.users = {
            0: self.create_user("school", 0, tel="13800138000", leader="张三"),
            1: self.create_user("city", 1),
            2: self.create_user("committee", 2),
            3: self.create_user("admin", 3),
            4: self.create_user("province", 4),
        }

    def test_every_role_reads_its_own_record(self):
        for user_type, user in self.users.items():
            with self.subTest(user_type=user_type):
                self.authorize_as(user)
                result = self.client.get("/api/user")
                self.assertEqual(result.status_code, 200)
                payload = result.json()
                self.assertEqual(payload["code"], 0)
                self.assertEqual(payload["msg"], "获取成功")
                self.assertEqual(payload["data"]["id"], user.id)
                self.assertEqual(payload["data"]["username"], user.username)
                self.assertNotIn("password", payload["data"])

    def test_optional_fields_are_normalized_to_empty_strings(self):
        user = self.users[0]
        self.authorize_as(user)
        payload = self.client.get("/api/user").json()["data"]
        for field in ("nickname", "description", "tel", "leader"):
            self.assertIn(field, payload)
            self.assertNotEqual(payload[field], None)
        self.assertEqual(payload["tel"], "13800138000")
        self.assertEqual(payload["leader"], "张三")

    def test_get_user_requires_token(self):
        self.assertEqual(self.client.get("/api/user").status_code, 401)

    def test_put_then_get_roundtrip(self):
        user = self.users[0]
        self.authorize_as(user)
        update = self.json_request("put", "/api/user",
                                   {"id": user.id, "nickname": "新昵称"})
        self.assertEqual(update.json()["code"], 0)
        payload = self.client.get("/api/user").json()["data"]
        self.assertEqual(payload["nickname"], "新昵称")
