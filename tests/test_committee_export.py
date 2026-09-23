"""组委会整组导出接口 + /api/export/data 数据隔离回归。

背景：原版 `Api\\ExportController::exportReportData`（`GET /api/export/data`）只按可选
`group` 过滤、不区分身份，7 个组委会页面都借它做整组导出。Django 为修 P1-2 跨校数据
泄漏给该接口加了 `filter(user_id=request.auth.id)`（提交 07056dac），学校侧隔离生效，
但组委会名下没有 Report，导出只剩表头。因此新增 `GET /api/committee/export/data`
（`role_error(request, 2)` + 只按 `group` 过滤），把两种数据范围拆开。

本文件同时锁住两件事：组委会能导整组；普通账号仍然只能导自己名下的。
"""
import io
from email.header import decode_header

from openpyxl import load_workbook

from apps.api.export_services import REPORT_DATA_HEADINGS, report_data_rows
from apps.core.models import Report

from .base import ApiTestCase

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def disposition_filename(response):
    """Django RFC2047-encodes non-ASCII header values; decode before asserting."""

    parts = decode_header(response["Content-Disposition"])
    return "".join(
        value.decode(encoding or "utf-8") if isinstance(value, bytes) else value
        for value, encoding in parts
    )


def sheet_rows(response):
    """Read an xlsx response body back as a list of rows (None → "")."""

    workbook = load_workbook(io.BytesIO(response.content))
    return [
        ["" if value is None else value for value in row]
        for row in workbook.active.iter_rows(values_only=True)
    ]


class CommitteeExportTests(ApiTestCase):
    def setUp(self):
        self.school_a = self.create_user("school-a", 0, nickname="A校")
        self.school_b = self.create_user("school-b", 0, nickname="B校")
        self.committee = self.create_user("committee", 2)
        self.city = self.create_user("city", 1)
        self.admin = self.create_user("admin", 3)

        self.a_middle = self.make_report(self.school_a, group="中学组", choir_name="A校中学团")
        self.b_middle = self.make_report(self.school_b, group="中学组", choir_name="B校中学团")
        self.b_primary = self.make_report(self.school_b, group="小学组", choir_name="B校小学团")

    # ---- 1. 组委会：按组导出整组 ----

    def test_committee_exports_every_school_in_the_group(self):
        self.authorize_as(self.committee)
        response = self.client.get("/api/committee/export/data?group=中学组")

        self.assertEqual(response.status_code, 200)
        rows = sheet_rows(response)
        self.assertEqual(rows[0], REPORT_DATA_HEADINGS)
        self.assertEqual([row[1] for row in rows[1:]], ["A校中学团", "B校中学团"])
        # 归属单位取自各报名自己的账号昵称，证明跨校数据确实都在
        self.assertEqual([row[0] for row in rows[1:]], ["A校", "B校"])

    def test_committee_export_is_not_limited_to_its_own_reports(self):
        """组委会名下没有任何 Report，仍应导得出数据（本次故障的直接回归）。"""

        self.assertEqual(Report.objects.filter(user_id=self.committee.id).count(), 0)
        self.authorize_as(self.committee)
        rows = sheet_rows(self.client.get("/api/committee/export/data?group=中学组"))

        self.assertEqual(len(rows) - 1, 2)

    def test_group_filter_is_applied_and_ordered_by_id(self):
        self.authorize_as(self.committee)

        primary = sheet_rows(self.client.get("/api/committee/export/data?group=小学组"))
        self.assertEqual([row[1] for row in primary[1:]], ["B校小学团"])

        # 无数据的组别：只有表头
        empty = sheet_rows(self.client.get("/api/committee/export/data?group=大学组"))
        self.assertEqual(len(empty), 1)

        # 不带 group：全部按 id 升序
        everything = sheet_rows(self.client.get("/api/committee/export/data"))
        self.assertEqual(
            [row[1] for row in everything[1:]],
            ["A校中学团", "B校中学团", "B校小学团"],
        )

    def test_every_chinese_group_value_exports_its_own_rows(self):
        """原版使用的三种中文组别（小学组/中学组/大学组）各自都能导全组。"""

        university = self.make_report(self.school_b, group="大学组", choir_name="B校大学团")
        self.authorize_as(self.committee)

        for group, expected in (
            ("小学组", ["B校小学团"]),
            ("中学组", ["A校中学团", "B校中学团"]),
            ("大学组", [university.choir_name]),
        ):
            with self.subTest(group=group):
                rows = sheet_rows(self.client.get("/api/committee/export/data", {"group": group}))
                self.assertEqual([row[1] for row in rows[1:]], expected)

    def test_committee_export_rows_match_the_plain_export_format(self):
        """列数、列顺序与行内容都与 /api/export/data 共用同一套 19 列。"""

        self.authorize_as(self.committee)
        rows = sheet_rows(self.client.get("/api/committee/export/data?group=中学组"))

        self.assertEqual(len(rows[0]), 19)
        self.assertEqual(rows[1], report_data_rows([self.a_middle])[1])

    # ---- 2. 非组委会账号：403 ----

    def test_other_roles_cannot_use_the_committee_endpoint(self):
        for user in (self.school_a, self.city, self.admin):
            with self.subTest(role=user.type):
                self.authorize_as(user)
                response = self.client.get("/api/committee/export/data?group=中学组")
                self.assertEqual(response.status_code, 403)
                self.assertEqual(response.json()["code"], "FORBIDDEN")

    def test_anonymous_is_rejected(self):
        self.clear_authorization()
        response = self.client.get("/api/committee/export/data?group=中学组")

        self.assertEqual(response.status_code, 401)

    # ---- 3. 普通导出接口的数据隔离（P1-2 回归） ----

    def test_plain_export_still_returns_only_the_callers_own_reports(self):
        self.authorize_as(self.school_a)
        response = self.client.get("/api/export/data")

        self.assertEqual(response.status_code, 200)
        rows = sheet_rows(response)
        self.assertEqual(rows[0], REPORT_DATA_HEADINGS)
        self.assertEqual([row[1] for row in rows[1:]], ["A校中学团"])
        self.assertNotIn("B校中学团", str(rows))
        self.assertNotIn("B校小学团", str(rows))

    def test_plain_export_group_filter_still_scoped_to_the_caller(self):
        """学校 B 有小学组数据，学校 A 用同样的 group 也必须导不到。"""

        self.authorize_as(self.school_a)
        rows = sheet_rows(self.client.get("/api/export/data?group=小学组"))

        self.assertEqual(len(rows), 1)

    def test_plain_export_is_unchanged_for_committee_accounts(self):
        """组委会走普通导出仍按本人过滤（0 行），整组数据只从新接口出去。"""

        self.authorize_as(self.committee)
        rows = sheet_rows(self.client.get("/api/export/data?group=中学组"))

        self.assertEqual(len(rows), 1)

    # ---- 4. 响应形态 ----

    def test_committee_export_returns_a_readable_xlsx(self):
        self.authorize_as(self.committee)
        response = self.client.get("/api/committee/export/data?group=中学组")

        self.assertEqual(response["Content-Type"], XLSX_CONTENT_TYPE)
        # 与原版 ExportController::exportReportData 的文件名一致
        self.assertEqual(disposition_filename(response), 'attachment; filename="数据导出.xlsx"')
        self.assertEqual(len(sheet_rows(response)), 3)
